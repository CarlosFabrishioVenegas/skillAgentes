#!/usr/bin/env python3
"""Convierte AK_kickoff_presentacion.md en un .pptx de kick-off (AK · VIRF).

Uso:
    python kickoff_md_a_pptx.py AK_kickoff_presentacion.md --salida AK_kickoff.pptx [--plantilla plantilla.pptx]
"""
import argparse
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

RE_CITA = re.compile(r"\s*\(([^()]*(?:\.md|\.pdf|\.docx|\.xlsx|sesión)[^()]*)\)")
AZUL, GRIS, ROJO = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x55, 0x55, 0x55), RGBColor(0xB0, 0x2A, 0x1E)


def secciones(md):
    out, actual = {}, "_titulo"
    for l in md.splitlines():
        if l.startswith("### "):
            actual = l[4:].strip()
        out.setdefault(actual, []).append(l)
    return out


def extraer_citas(texto, citas, usadas=None):
    def rep(m):
        c = m.group(1).strip()
        if c not in citas:
            citas.append(c)
        if usadas is not None:
            usadas.add(citas.index(c) + 1)
        return f" [{citas.index(c) + 1}]"
    return RE_CITA.sub(rep, texto).strip()


def caja_texto(slide, x, y, w, h, parrafos, tam=14, negrita_primero=False, color=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, p in enumerate(parrafos):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.text = p
        par.font.size = Pt(tam)
        if color:
            par.font.color.rgb = color
        if negrita_primero and i == 0:
            par.font.bold = True
    # reducción simple por longitud
    largo = sum(len(p) for p in parrafos)
    if largo > 900:
        for par in tf.paragraphs:
            par.font.size = Pt(max(11, tam - 3))
    return tb


def pie(slide, prs, citas_slide):
    if not citas_slide:
        return
    txt = "  ".join(f"[{i}] {c}" for i, c in citas_slide)
    caja_texto(slide, Inches(0.4), prs.slide_height - Inches(0.75), prs.slide_width - Inches(0.8), Inches(0.6), [txt], tam=9, color=GRIS)


def titulo(slide, texto, prs):
    if slide.shapes.title is not None:
        t = slide.shapes.title
        t.left, t.top, t.width, t.height = Inches(0.4), Inches(0.3), prs.slide_width - Inches(0.8), Inches(0.9)
        t.text = texto
    else:
        caja_texto(slide, Inches(0.4), Inches(0.3), prs.slide_width - Inches(0.8), Inches(0.8), [texto], tam=26, color=AZUL)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("md")
    ap.add_argument("--salida", default="AK_kickoff.pptx")
    ap.add_argument("--plantilla", default=None)
    a = ap.parse_args()
    md = open(a.md, encoding="utf-8").read()
    sec = secciones(md)
    prs = Presentation(a.plantilla) if a.plantilla else Presentation()
    if not a.plantilla:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    lay_titulo = prs.slide_layouts[0]
    lay_blanco = prs.slide_layouts[5] if len(prs.slide_layouts) > 5 else prs.slide_layouts[-1]
    titulo_md = next((l[3:].strip() for l in md.splitlines() if l.startswith("## ")), "Kick-off — No evidenciado")
    parcial = next((l.lstrip("> ").strip() for l in md.splitlines() if "parcial" in l.lower() and l.startswith(">")), None)

    # Portada
    s = prs.slides.add_slide(lay_titulo)
    s.shapes.title.text = titulo_md
    if not a.plantilla:
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Inches(0.8), Inches(2.3), prs.slide_width - Inches(1.6), Inches(1.6)
    if len(s.placeholders) > 1:
        s.placeholders[1].text = "Validación periódica · Kick-off de entendimiento"
        if not a.plantilla:
            st = s.placeholders[1]
            st.left, st.top, st.width, st.height = Inches(0.8), Inches(4.1), prs.slide_width - Inches(1.6), Inches(1.0)
    if parcial:
        caja_texto(s, Inches(0.4), Inches(0.1), prs.slide_width - Inches(0.8), Inches(0.5), [parcial], tam=12, color=ROJO)

    citas_global = []

    # Lámina 1
    l1 = next((v for k, v in sec.items() if k.lower().startswith("lámina 1") or k.lower().startswith("lamina 1")), None)
    if l1 is None:
        print("[AVISO] Falta 'Lámina 1' en el Markdown: se genera con 'No evidenciado'")
        l1 = ["**Alcance**", "- No evidenciado", "**Limitaciones**", "- No evidenciado"]
    bloques, actual = {"Alcance": [], "Limitaciones": []}, None
    for l in l1:
        m = re.match(r"\*\*(Alcance|Limitaciones)\*\*", l.strip())
        if m:
            actual = m.group(1); continue
        if actual and l.strip().startswith("-"):
            bloques[actual].append(l.strip()[1:].strip())
    s = prs.slides.add_slide(lay_blanco)
    titulo(s, "Alcance y Limitaciones", prs)
    usadas = set()
    alc = [extraer_citas("• " + x, citas_global, usadas) for x in bloques["Alcance"]] or ["• No evidenciado"]
    lim = [extraer_citas("• " + x, citas_global, usadas) for x in bloques["Limitaciones"]] or ["• No evidenciado"]
    w = (prs.slide_width - Inches(1.2)) / 2
    caja_texto(s, Inches(0.4), Inches(1.4), w, Inches(5.2), ["Alcance"] + alc, tam=15, negrita_primero=True)
    caja_texto(s, Inches(0.8) + w, Inches(1.4), w, Inches(5.2), ["Limitaciones"] + lim, tam=15, negrita_primero=True)
    pie(s, prs, [(i, citas_global[i - 1]) for i in sorted(usadas)])

    # Lámina 2
    l2 = next((v for k, v in sec.items() if k.lower().startswith("lámina 2") or k.lower().startswith("lamina 2")), [])
    fila = None
    for l in l2:
        if l.strip().startswith("|") and "→" in l and "Input de Datos" not in l and not re.match(r"^\|\s*-", l.strip()):
            fila = [c.strip() for c in l.strip().strip("|").split("|")]
    if not fila:
        print("[AVISO] Falta la tabla del flujo: cajas con 'No evidenciado'")
        fila = ["No evidenciado", "→", "No evidenciado", "→", "No evidenciado"]
    cajas = [fila[0], fila[2] if len(fila) > 2 else "No evidenciado", fila[4] if len(fila) > 4 else "No evidenciado"]
    notas = [l.strip() for l in l2 if l.strip().startswith("**Notas")]
    s = prs.slides.add_slide(lay_blanco)
    titulo(s, "Flujo de datos", prs)
    usadas = set()
    nombres = ["Input de Datos", "Tratamiento y resultado de los datos", "Usuarios Finales"]
    ancho = (prs.slide_width - Inches(1.6)) / 3
    for i, (n, c) in enumerate(zip(nombres, cajas)):
        x = Inches(0.4) + i * (ancho + Inches(0.4))
        rect = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.4), ancho, Inches(4.2))
        rect.fill.solid(); rect.fill.fore_color.rgb = RGBColor(0xEE, 0xF2, 0xF7); rect.line.color.rgb = AZUL
        caja_texto(s, x + Inches(0.1), Inches(1.5), ancho - Inches(0.2), Inches(4.0),
                   [n, extraer_citas(c, citas_global, usadas)], tam=14, negrita_primero=True)
        if i < 2:
            fl = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + ancho + Inches(0.05), Inches(3.3), Inches(0.3), Inches(0.4))
            fl.fill.solid(); fl.fill.fore_color.rgb = AZUL
    if notas:
        caja_texto(s, Inches(0.4), Inches(5.75), prs.slide_width - Inches(0.8), Inches(0.9),
                   [extraer_citas(re.sub(r"\*\*", "", n), citas_global, usadas) for n in notas], tam=10)
    pie(s, prs, [(i, citas_global[i - 1]) for i in sorted(usadas)])

    # Fuentes
    fu = next((v for k, v in sec.items() if k.lower().startswith("fuentes")), ["- No evidenciado"])
    s = prs.slides.add_slide(lay_blanco)
    titulo(s, "Fuentes y consultas abiertas", prs)
    lineas = [l.strip() for l in fu[1:] if l.strip()]
    lineas += ["", "Referencias de las láminas:"] + [f"[{i+1}] {c}" for i, c in enumerate(citas_global)]
    caja_texto(s, Inches(0.4), Inches(1.4), prs.slide_width - Inches(0.8), Inches(5.6), lineas, tam=12)

    prs.save(a.salida)
    print(f"Presentación generada: {a.salida} · {len(prs.slides)} láminas · {len(citas_global)} citas")


if __name__ == "__main__":
    main()
