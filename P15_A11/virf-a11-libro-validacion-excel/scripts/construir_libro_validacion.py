#!/usr/bin/env python3
"""Construye Validacion_<codigo>_<ciclo>.xlsx con las 7 hojas + Observaciones (A11 · VIRF).

Uso:
    python construir_libro_validacion.py --codigo MOD-0412 --ciclo 2026 --consolidado consolidado.csv \
        --pruebas pruebas_a2.csv pruebas_a7.csv --evidencias evidencias_hash.csv --plan plan_accion.csv \
        --conclusion conclusion.md --salida Validacion_MOD-0412_2026.xlsx [--markdown]
"""
import argparse
import glob
import re
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

COLORES = {"Cumple": "C6EFCE", "Cumple con observación": "FFEB9C", "No cumple": "FFC7CE",
           "No evidenciado": "D9D9D9", "No verificable": "D9D9D9", "No aplica": "FFFFFF",
           "PASA": "C6EFCE", "FALLA": "FFC7CE"}
CONTROLES = ["C1", "C2", "C3", "C4", "C5", "C6"]


def leer_varios(patrones):
    fs = [f for p in patrones for f in glob.glob(p)]
    return pd.concat([pd.read_csv(f, dtype=str) for f in fs], ignore_index=True) if fs else pd.DataFrame()


def semaforo(estados):
    e = set(estados)
    if "No cumple" in e:
        return "Rojo"
    if e & {"No evidenciado", "No verificable"}:
        return "Gris"
    if "Cumple con observación" in e:
        return "Ámbar"
    return "Verde" if e else "Sin filas"


def veredicto(C, cobertura):
    crit_impl_uso = C[(C.estado == "No cumple") & (C.severidad.str.lower() == "crítica") & (C.control.isin(["C4", "C6"]))]
    if cobertura < 100:
        return "Validación incompleta"
    if len(crit_impl_uso):
        return "Validación no satisfactoria"
    if (C.estado == "No cumple").any() or (C.estado == "Cumple con observación").any():
        return "Validación con observaciones"
    return "Validación satisfactoria"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--codigo", required=True); ap.add_argument("--ciclo", required=True)
    ap.add_argument("--consolidado", required=True)
    ap.add_argument("--pruebas", nargs="*", default=[])
    ap.add_argument("--evidencias", default=None)
    ap.add_argument("--plan", default=None)
    ap.add_argument("--conclusion", default=None)
    ap.add_argument("--salida", default=None)
    ap.add_argument("--markdown", action="store_true")
    a = ap.parse_args()
    salida = a.salida or f"Validacion_{a.codigo}_{a.ciclo}.xlsx"

    C = pd.read_csv(a.consolidado, dtype=str).fillna("")
    for c in ("severidad", "accion_pendiente", "responsable", "fecha_comprometida", "id_pruebas", "evidencias", "evaluacion_sustentada"):
        if c not in C:
            C[c] = ""
    P = leer_varios(a.pruebas).fillna("")
    E = pd.read_csv(a.evidencias, dtype=str).fillna("") if a.evidencias else pd.DataFrame(columns=["archivo", "tipo", "sha256", "fecha_modificacion", "controles"])
    PL = pd.read_csv(a.plan, dtype=str).fillna("") if a.plan else pd.DataFrame(columns=["accion", "responsable", "fecha", "control_origen", "severidad"])
    concl = open(a.conclusion, encoding="utf-8").read().strip().splitlines() if a.conclusion else ["(pendiente)"]
    if len([l for l in concl if l.strip()]) > 12:
        print("[AVISO] La conclusión supera 12 líneas.")

    # verificaciones
    alertas = []
    ids_p = set(P["ID prueba"]) if "ID prueba" in P else set()
    for i, r in C.iterrows():
        if r.tipo == "Cuantitativo" and r.estado == "Cumple":
            ids = [x.strip() for x in re.split(r"[;,]", r.id_pruebas) if x.strip()]
            if not ids or not set(ids) <= ids_p:
                alertas.append(f"{r.id_virf}: control cuantitativo 'Cumple' sin prueba en HOJA 3")
        if not r.fuente_exacta:
            alertas.append(f"{r.id_virf}: fuente vacía")
    for i, r in PL.iterrows():
        if not r.get("responsable") or not r.get("fecha"):
            alertas.append(f"Plan de acción fila {i+2}: sin responsable o fecha")

    # cobertura por actividad única (una actividad puede tener filas de referencia cruzada en otros controles)
    U = C.assign(_ev=~C.estado.isin(["No evidenciado"])).groupby("id_virf").agg(control=("control", "first"), _ev=("_ev", "any"))
    aplicables, evidenciadas = len(U), int(U._ev.sum())
    cob_glob = 100 * evidenciadas / max(aplicables, 1)
    res = [{"Indicador": "Veredicto global", "Valor": veredicto(C, cob_glob)},
           {"Indicador": "Cobertura global (% actividades evidenciadas)", "Valor": f"{cob_glob:.1f} %"}]
    for c in CONTROLES:
        sub = C[C.control == c]
        su = U[U.control == c]
        cob = 100 * su._ev.sum() / max(len(su), 1)
        res.append({"Indicador": f"{c} · semáforo / cobertura", "Valor": f"{semaforo(sub.estado)} / {cob:.1f} % ({len(su)} actividades)"})
    res += [{"Indicador": "Alertas de integridad", "Valor": len(alertas)}] + [{"Indicador": "·", "Valor": x} for x in alertas[:50]]

    h2 = C.rename(columns={"id_virf": "ID VIRF", "control": "Ámbito (C1–C6)", "actividad": "Actividad", "estado": "Estado",
                           "resultado": "Resultado", "evaluacion_sustentada": "Evaluación sustentada", "fuente_exacta": "Fuente exacta",
                           "severidad": "Severidad", "accion_pendiente": "Acción pendiente", "responsable": "Responsable",
                           "fecha_comprometida": "Fecha comprometida", "id_pruebas": "ID prueba(s)", "evidencias": "Evidencia(s)",
                           "agente_origen": "Agente de origen", "tipo": "Cuantitativo/Cualitativo"})
    cols2 = ["ID VIRF", "Ámbito (C1–C6)", "Actividad", "Estado", "Resultado", "Evaluación sustentada", "Fuente exacta", "Severidad",
             "Acción pendiente", "Responsable", "Fecha comprometida", "ID prueba(s)", "Evidencia(s)", "Agente de origen", "Cuantitativo/Cualitativo"]
    h2 = h2[[c for c in cols2 if c in h2]]
    h4 = E.rename(columns={"archivo": "Archivo", "tipo": "Tipo", "sha256": "Hash SHA-256", "fecha_modificacion": "Fecha", "controles": "Control(es) que sustenta"})
    br = C[C.estado.isin(["No evidenciado", "No verificable"])]
    h5 = pd.DataFrame({"Ítem": br.id_virf + " · " + br.actividad, "Control": br.control, "Estado": br.estado,
                       "Evidencia faltante": br.resultado, "Impacto": br.severidad, "Acción de subsanación": br.accion_pendiente})
    h6 = PL.rename(columns={"accion": "Acción", "responsable": "Responsable", "fecha": "Fecha", "control_origen": "Control de origen", "severidad": "Severidad"})
    if "Fecha" in h6:
        h6 = h6.sort_values("Fecha")
    h7 = pd.DataFrame({"Conclusión": concl + ["", "Firma del validador: ____________________   Fecha: ________",
                                              "Firma del revisor:   ____________________   Fecha: ________"]})
    obs = C[C.estado.isin(["No cumple", "Cumple con observación", "No evidenciado", "No verificable"])]
    h8 = pd.DataFrame({"Ámbito": obs.control, "Riesgo": obs.severidad, "Sustentación": obs.fuente_exacta,
                       "Observaciones": obs.evaluacion_sustentada.where(obs.evaluacion_sustentada != "", obs.resultado),
                       "Recomendaciones": obs.accion_pendiente, "Responsable/Fecha": obs.responsable + " / " + obs.fecha_comprometida})
    hojas = [("Resumen", pd.DataFrame(res)), ("Controles", h2), ("Pruebas", P), ("Evidencia", h4), ("Brechas", h5),
             ("Plan de acción", h6), ("Conclusión", h7), ("Observaciones", h8)]
    with pd.ExcelWriter(salida, engine="openpyxl") as xw:
        for n, df in hojas:
            (df if len(df.columns) else pd.DataFrame({"(sin datos)": []})).to_excel(xw, sheet_name=n, index=False)
    wb = load_workbook(salida)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for col in ws.columns:
            largo = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max(12, largo * 0.9), 60)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if str(c.value) in COLORES:
                    c.fill = PatternFill("solid", fgColor=COLORES[str(c.value)])
    wb.save(salida)
    print(f"Libro generado: {salida} · veredicto: {res[0]['Valor']} · cobertura {cob_glob:.1f} % · alertas {len(alertas)}")
    for x in alertas[:20]:
        print("  [ALERTA]", x)
    if a.markdown:
        md = [f"Copiar cada bloque HOJA N a una pestaña homónima del libro Validacion_{a.codigo}.xlsx", ""]
        for i, (n, df) in enumerate(hojas, 1):
            md += [f"=== HOJA {i}: {n} ===", df.to_markdown(index=False) if len(df) else "(sin filas)", ""]
        open(salida.replace(".xlsx", "_hojas.md"), "w", encoding="utf-8").write("\n".join(md))
        print("Versión Markdown:", salida.replace(".xlsx", "_hojas.md"))


if __name__ == "__main__":
    main()
