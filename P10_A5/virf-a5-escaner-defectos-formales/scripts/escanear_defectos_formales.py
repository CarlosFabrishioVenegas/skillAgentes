#!/usr/bin/env python3
"""Escanea manuales PDF/Word en busca de defectos formales de auditabilidad (A5 · VIRF).

Uso:
    python escanear_defectos_formales.py <carpeta_o_archivo> [--entregados <carpeta>] --salida defectos.csv
"""
import argparse
import csv
import re
import sys
from pathlib import Path

RELLENO = re.compile(r"(\bpor definir\b|\bpendiente\b|\bTBD\b|\bTODO\b|\bPOR COMPLETAR\b|[\[<]\s*completar|\bxxx+\b|\[insertar|\ba definir\b|\ben construcci[oó]n\b)")
RELLENO_I = re.compile(r"(\bpor definir\b|\bpendiente\b|\ba definir\b|\ben construcci[oó]n\b)", re.I)
VERSION = re.compile(r"(versi[oó]n\s*[:\-]?\s*\d+(\.\d+)*|\bv\s?\d+(\.\d+)+\b)", re.I)
FECHA = re.compile(r"(\b\d{1,2}[/\-.]\d{1,2}[/\-.]\d{2,4}\b|\b(19|20)\d{2}[/\-.]\d{1,2}[/\-.]\d{1,2}\b|\b\d{1,2} de [a-záéíóú]+ de (19|20)\d{2}\b|\b[a-záéíóú]+ (de )?(19|20)\d{2}\b)", re.I)
AUTOR = re.compile(r"(elaborado por|autor|aprobado por|revisado por|responsable del documento)", re.I)
BORRADOR = re.compile(r"\b(borrador|draft|versi[oó]n preliminar)\b", re.I)
ANEXO = re.compile(r"\banexo\s+([A-Z0-9]{1,3})\b", re.I)
FIG = re.compile(r"\b(figura|gr[aá]fico|tabla|cuadro)\s+(\d+(\.\d+)?)", re.I)
FUENTE = re.compile(r"\bfuente\s*:", re.I)
CAMBIOS = re.compile(r"(control de cambios|historial de versiones|registro de cambios|control de versiones)", re.I)
TITULO = re.compile(r"^\s*\d+(\.\d+){0,3}\.?\s+[A-ZÁÉÍÓÚÑ].{2,90}$")
INDICE = re.compile(r"(\.\s?){4,}|\s\d{1,3}\s*$")  # líneas de índice con puntos guía o número de página final
EXCLUIR_RELLENO = re.compile(r"acci[oó]n(es)? pendiente|pendiente de la curva|pendientes? de (pago|cobro)", re.I)


def paginas_pdf(p):
    import pdfplumber
    with pdfplumber.open(p) as pdf:
        for i, pg in enumerate(pdf.pages, 1):
            txt = pg.extract_text() or ""
            area_img = sum((im["x1"] - im["x0"]) * (im["bottom"] - im["top"]) for im in pg.images)
            yield str(i), txt, area_img / (pg.width * pg.height) if pg.width else 0


def paginas_docx(p):
    import docx
    d = docx.Document(p)
    bloque, n = [], 0
    for i, par in enumerate(d.paragraphs, 1):
        bloque.append(par.text)
        if len(bloque) >= 40:
            yield f"§{i-39}-{i}", "\n".join(bloque), 0; bloque = []
    for t in d.tables:
        bloque.append("\n".join(" | ".join(c.text for c in r.cells) for r in t.rows))
    if bloque:
        yield f"§fin", "\n".join(bloque), 0


def frag(txt, m):
    a = max(0, m.start() - 50)
    return txt[a:m.end() + 60].replace("\n", " ")[:120]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("ruta")
    ap.add_argument("--entregados", default=None)
    ap.add_argument("--salida", default="defectos_formales.csv")
    a = ap.parse_args()
    base = Path(a.ruta)
    docs = [base] if base.is_file() else sorted(x for x in base.rglob("*") if x.suffix.lower() in (".pdf", ".docx"))
    entregados = [x.name.lower() for x in Path(a.entregados or a.ruta).rglob("*")] if (a.entregados or base.is_dir()) else []
    filas = []

    def add(doc, pag, cod, defecto, evid, sev):
        filas.append({"archivo": doc.name, "pagina": pag, "codigo": cod, "defecto": defecto, "evidencia": evid, "severidad_sugerida": sev})

    for doc in docs:
        gen = paginas_pdf(doc) if doc.suffix.lower() == ".pdf" else paginas_docx(doc)
        paginas = list(gen)
        inicio = "\n".join(t for _, t, _ in paginas[:3])
        todo = "\n".join(t for _, t, _ in paginas)
        if not VERSION.search(inicio):
            add(doc, "1-3", "F03", "Sin versión en portada", "", "Media (Tier 3) / No cumple (Tier 1-2)")
        if not FECHA.search(inicio):
            add(doc, "1-3", "F04", "Sin fecha en portada", "", "Media (Tier 3) / No cumple (Tier 1-2)")
        if not AUTOR.search(todo):
            add(doc, "-", "F05", "Sin autor/aprobador identificable", "", "Media")
        if not CAMBIOS.search(todo):
            add(doc, "-", "F11", "Sin control de cambios / historial de versiones", "", "Media")
        anexos_vistos = set()
        for pag, txt, prop_img in paginas:
            for m in list(RELLENO.finditer(txt)) + [x for x in RELLENO_I.finditer(txt) if not RELLENO.match(txt, x.start())]:
                if EXCLUIR_RELLENO.search(frag(txt, m)):
                    continue
                add(doc, pag, "F01", "Texto de relleno", frag(txt, m), "Media")
            for m in BORRADOR.finditer(txt):
                add(doc, pag, "F06", "Marca de borrador", frag(txt, m), "Alta"); break
            lineas = [l for l in txt.splitlines() if l.strip()]
            for l1, l2 in zip(lineas, lineas[1:]):
                es_tit = lambda l: TITULO.match(l) and not INDICE.search(l) and not l.rstrip().endswith((".", ":", ";", ",")) and len(l.split()) <= 12
                if es_tit(l1) and es_tit(l2):
                    add(doc, pag, "F02", "Sección posiblemente vacía", f"{l1.strip()} → {l2.strip()}"[:120], "Media")
            for m in ANEXO.finditer(txt):
                anexos_vistos.add((m.group(1).upper(), pag))
            if FIG.search(txt) and not FUENTE.search(txt):
                add(doc, pag, "F08", "Figura/tabla sin fuente en la página", frag(txt, FIG.search(txt)), "Baja")
            if prop_img > 0.45 and len(txt.strip()) < 200:
                add(doc, pag, "F09", "Página dominada por imagen/captura", f"{prop_img:.0%} del área", "Observación")
        for anx, pag in sorted(anexos_vistos):
            if entregados and not any(re.search(rf"anexo[\s_\-]*{re.escape(anx.lower())}\b", n) for n in entregados):
                add(doc, pag, "F07", f"Anexo {anx} referenciado y no entregado", f"Anexo {anx}", "Alta")
    # F10 versiones duplicadas
    raiz = {}
    for d in docs:
        k = re.sub(r"(_?v\d+(\.\d+)*|_?\d{4}[-_]?\d{2}([-_]?\d{2})?|_?(final|ult|vf))", "", d.stem.lower())
        raiz.setdefault(k, []).append(d.name)
    for k, ns in raiz.items():
        if len(ns) > 1:
            filas.append({"archivo": " | ".join(ns), "pagina": "-", "codigo": "F10", "defecto": "Varias versiones del mismo documento",
                          "evidencia": k, "severidad_sugerida": "Media"})
    if not docs:
        sys.exit("No se encontraron PDF/DOCX.")
    with open(a.salida, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["archivo", "pagina", "codigo", "defecto", "evidencia", "severidad_sugerida"])
        w.writeheader(); w.writerows(filas)
    print(f"{len(docs)} documentos · {len(filas)} defectos → {a.salida}")
    cont = {}
    for r in filas:
        cont[r["codigo"]] = cont.get(r["codigo"], 0) + 1
    for k in sorted(cont):
        print(f"  {k}: {cont[k]}")


if __name__ == "__main__":
    main()
