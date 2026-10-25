#!/usr/bin/env python3
"""Inventario forense de un libro Excel de modelo (A2 · VIRF).

Uso:
    python inventario_excel.py libro.xlsx --salida inventario/
No modifica el libro (solo lectura).
"""
import argparse
import csv
import hashlib
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl: pip install openpyxl")

VOLATILES = {"TODAY", "NOW", "RAND", "RANDBETWEEN", "OFFSET", "INDIRECT", "HOY", "AHORA", "ALEATORIO", "DESREF", "INDIRECTO"}
RE_FUNC = re.compile(r"([A-Z][A-Z0-9\.]+)\(")
RE_REF_HOJA = re.compile(r"(?:'([^']+)'|([A-Za-z_][\w\.]*))!")
RE_RANGO_GRANDE = re.compile(r"\$?[A-Z]{1,3}\$?(\d+):\$?[A-Z]{1,3}\$?(\d+)")


def escribir(ruta, filas):
    if not filas:
        Path(ruta).write_text("(sin registros)\n", encoding="utf-8"); return
    with open(ruta, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(filas[0].keys())); w.writeheader(); w.writerows(filas)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("libro")
    ap.add_argument("--salida", default="inventario")
    a = ap.parse_args()
    ruta = Path(a.libro); out = Path(a.salida); out.mkdir(parents=True, exist_ok=True)
    sha = hashlib.sha256(ruta.read_bytes()).hexdigest()
    alertas = []

    with zipfile.ZipFile(ruta) as z:
        nombres = z.namelist()
        if any("vbaProject.bin" in n for n in nombres):
            alertas.append({"severidad": "Alta", "tipo": "VBA", "ubicacion": "vbaProject.bin", "detalle": "Macro VBA presente: leer con olevba o declarar 'Lógica en VBA no accesible'"})
        for n in nombres:
            if n.startswith("xl/externalLinks/") and n.endswith(".xml"):
                alertas.append({"severidad": "Alta", "tipo": "Vínculo externo", "ubicacion": n, "detalle": "Referencia a libro externo no entregado"})
        if "xl/connections.xml" in nombres:
            alertas.append({"severidad": "Alta", "tipo": "Conexión de datos", "ubicacion": "xl/connections.xml", "detalle": "Conexión a fuente externa (ODBC/Power Query)"})
        wbxml = z.read("xl/workbook.xml").decode("utf-8", "replace")
        if 'iterate="1"' in wbxml or "iterate=\"true\"" in wbxml:
            alertas.append({"severidad": "Alta", "tipo": "Cálculo iterativo", "ubicacion": "workbook.xml", "detalle": "Iteración habilitada: replicar n° de iteraciones"})

    wb_f = openpyxl.load_workbook(ruta, data_only=False, keep_vba=False)
    wb_v = openpyxl.load_workbook(ruta, data_only=True, keep_vba=False)
    if getattr(wb_f, "epoch", None) and wb_f.epoch.year == 1904:
        alertas.append({"severidad": "Media", "tipo": "Sistema de fechas 1904", "ubicacion": "libro", "detalle": "Ajustar conversión de fechas"})

    hojas, formulas, nombres_rows = [], [], []
    for dn in wb_f.defined_names.values() if hasattr(wb_f.defined_names, "values") else wb_f.defined_names.definedName:
        nombres_rows.append({"nombre": dn.name, "destino": dn.attr_text})

    for ws in wb_f.worksheets:
        wsv = wb_v[ws.title]
        n_form = n_const = n_num = 0
        sin_cache = 0
        por_col = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                es_f = isinstance(v, str) and v.startswith("=")
                col = c.column_letter
                por_col.setdefault(col, Counter())["f" if es_f else "c"] += 1
                if es_f:
                    n_form += 1
                    funcs = sorted(set(RE_FUNC.findall(v.upper())))
                    vol = bool(set(funcs) & VOLATILES)
                    refs = sorted({(m[0] or m[1]) for m in RE_REF_HOJA.findall(v)})
                    if wsv[c.coordinate].value is None:
                        sin_cache += 1
                    formulas.append({"hoja": ws.title, "celda": c.coordinate, "formula": v,
                                     "funciones": ";".join(funcs), "hojas_referenciadas": ";".join(refs),
                                     "volatil": "sí" if vol else "no", "valor_cacheado": wsv[c.coordinate].value})
                    if vol:
                        alertas.append({"severidad": "Media", "tipo": "Fórmula volátil", "ubicacion": f"{ws.title}!{c.coordinate}", "detalle": v[:120]})
                    for m in RE_RANGO_GRANDE.finditer(v):
                        if int(m.group(2)) > max(ws.max_row, 1) * 5 and int(m.group(2)) > 5000:
                            alertas.append({"severidad": "Baja", "tipo": "Rango sobredimensionado", "ubicacion": f"{ws.title}!{c.coordinate}", "detalle": m.group(0)})
                else:
                    n_const += 1
                    if isinstance(v, (int, float)):
                        n_num += 1
        # constantes numéricas en columnas mayoritariamente de fórmulas -> hardcoded
        for col, cnt in por_col.items():
            if cnt["f"] >= 10 and 0 < cnt["c"] <= cnt["f"] * 0.2:
                for row in ws.iter_rows(min_col=openpyxl.utils.column_index_from_string(col), max_col=openpyxl.utils.column_index_from_string(col)):
                    c = row[0]
                    if isinstance(c.value, (int, float)):
                        alertas.append({"severidad": "Alta", "tipo": "Cálculo endurecido", "ubicacion": f"{ws.title}!{c.coordinate}", "detalle": f"valor {c.value} en columna de fórmulas"})
        if ws.sheet_state != "visible":
            alertas.append({"severidad": "Media", "tipo": "Hoja oculta", "ubicacion": ws.title, "detalle": ws.sheet_state})
        if ws.auto_filter and ws.auto_filter.ref:
            alertas.append({"severidad": "Alta", "tipo": "Autofiltro activo", "ubicacion": f"{ws.title}!{ws.auto_filter.ref}", "detalle": "Verificar si el cálculo depende de filas visibles"})
        if sin_cache:
            alertas.append({"severidad": "Media", "tipo": "Fórmulas sin valor cacheado", "ubicacion": ws.title, "detalle": f"{sin_cache} celdas: recalcular y guardar en Excel antes de reconciliar"})
        total = n_form + n_const
        if total == 0:
            rol = "AUXILIAR (vacía)"
        elif n_form / total > 0.6:
            rol = "CÁLCULO o SALIDA"
        elif n_form == 0 and ws.max_row > 50:
            rol = "ENTRADA"
        elif n_form == 0:
            rol = "PARÁMETROS o AUXILIAR"
        else:
            rol = "MIXTA (revisar)"
        hojas.append({"hoja": ws.title, "visibilidad": ws.sheet_state, "dimensiones": ws.dimensions,
                      "n_formulas": n_form, "n_constantes": n_const, "n_constantes_numericas": n_num,
                      "rol_sugerido": rol, "filtro_activo": bool(ws.auto_filter and ws.auto_filter.ref)})

    escribir(out / "inventario_hojas.csv", hojas)
    escribir(out / "inventario_nombres.csv", nombres_rows)
    escribir(out / "formulas.csv", formulas)
    escribir(out / "alertas.csv", alertas)
    md = [f"# Inventario de {ruta.name}", f"SHA-256: `{sha}`", "",
          "| Hoja | Visibilidad | Fórmulas | Constantes | Rol sugerido | Filtro |", "|---|---|---|---|---|---|"]
    md += [f"| {h['hoja']} | {h['visibilidad']} | {h['n_formulas']} | {h['n_constantes']} | {h['rol_sugerido']} | {h['filtro_activo']} |" for h in hojas]
    md += ["", "## Alertas", "| Severidad | Tipo | Ubicación | Detalle |", "|---|---|---|---|"]
    md += [f"| {x['severidad']} | {x['tipo']} | {x['ubicacion']} | {x['detalle']} |" for x in sorted(alertas, key=lambda x: ["Alta", "Media", "Baja"].index(x["severidad"]))]
    (out / "resumen.md").write_text("\n".join(md), encoding="utf-8")
    print(f"{len(hojas)} hojas · {len(formulas)} fórmulas · {len(alertas)} alertas · sha256={sha}")


if __name__ == "__main__":
    main()
